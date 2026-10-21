import os
import io
import json
import secrets
import time
from typing import Optional, Any, List

from fastapi import FastAPI, Header, HTTPException, UploadFile, File, Depends, Request, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

import assistant
import llm
import rag
import store_db

app = FastAPI(title="OmniTag AI Server", description="OmniTag AR GenAI API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "omnitag-admin")
SESSION_TTL = 8 * 3600  # seconds
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
# Address phones use to reach this server (printed into QR codes), e.g. http://192.168.43.10:8000
PUBLIC_URL = os.environ.get("PUBLIC_URL", "").rstrip("/")
WEB_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "web")

# Slow down password guessing: max failed logins per client per window
LOGIN_MAX_FAILURES = 5
LOGIN_WINDOW = 300  # seconds

_sessions = {}  # token -> expiry timestamp
_failed_logins = {}  # client ip -> [timestamps]

store_db.init_db()
rag.start_background_index()


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

class TagRequest(BaseModel):
    product_name: str = ""
    qr_code: Optional[str] = None
    user_context: str = "General Buyer"
    product: Optional[dict] = None


class ChatMessage(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    product_name: str = ""
    qr_code: Optional[str] = None
    question: str = Field(min_length=1, max_length=2000)
    persona: str = "General Buyer"
    history: List[ChatMessage] = []
    product: Optional[dict] = None


@app.get("/api/health")
def health_check():
    return {
        "status": "online",
        "service": "FastAPI LLM Gateway",
        "backend": llm.OLLAMA_BASE,
        "model": llm.CHAT_MODEL,
        "rag": rag.status(),
    }


@app.get("/api/config")
def get_config():
    return {"public_url": PUBLIC_URL}


# ---------------------------------------------------------------------------
# Products (public, read-only)
# ---------------------------------------------------------------------------

@app.get("/api/products")
def get_products():
    return {"success": True, "products": store_db.list_products()}


@app.get("/api/products/{qr_code}")
def get_product(qr_code: str):
    product = store_db.get_product(qr_code)
    if not product:
        raise HTTPException(status_code=404, detail="Tag not found.")
    return {"success": True, "product": product}


@app.get("/api/search")
def search(q: str = Query(..., min_length=1, max_length=200), limit: int = 6):
    """Semantic product search with price/category filters parsed from the text."""
    q = q.strip()[:200]
    if not q:
        raise HTTPException(status_code=400, detail="Empty search.")
    try:
        found = rag.search_products(q, limit=max(1, min(limit, 12)))
    except (llm.LLMError, RuntimeError) as e:
        return {"success": False, "error": str(e)}
    return {"success": True, **found}


def resolve_product(qr_code, client_product):
    """Prefer the database record for LLM grounding; fall back to what the client sent."""
    if qr_code:
        product = store_db.get_product(qr_code)
        if product:
            return product
    return client_product


# ---------------------------------------------------------------------------
# Admin: login, price editing, Excel import/export
# ---------------------------------------------------------------------------

class LoginRequest(BaseModel):
    password: str


class PriceUpdate(BaseModel):
    price: Any = None
    mrp: Any = None
    amazon_price: Any = None
    flipkart_price: Any = None
    croma_price: Any = None


def require_admin(authorization: Optional[str] = Header(None)):
    token = (authorization or "").removeprefix("Bearer ").strip()
    expiry = _sessions.get(token)
    if not token or not expiry or expiry < time.time():
        _sessions.pop(token, None)
        raise HTTPException(status_code=401, detail="Please log in as admin.")
    return token


@app.post("/api/admin/login")
def admin_login(req: LoginRequest, request: Request):
    client = request.client.host if request.client else "unknown"
    now = time.time()
    recent = [t for t in _failed_logins.get(client, []) if now - t < LOGIN_WINDOW]
    if len(recent) >= LOGIN_MAX_FAILURES:
        raise HTTPException(status_code=429, detail="Too many failed attempts. Try again in a few minutes.")
    if not secrets.compare_digest(req.password.encode(), ADMIN_PASSWORD.encode()):
        _failed_logins[client] = recent + [now]
        raise HTTPException(status_code=401, detail="Wrong password.")
    _failed_logins.pop(client, None)
    token = secrets.token_urlsafe(32)
    _sessions[token] = time.time() + SESSION_TTL
    return {"success": True, "token": token, "expires_in": SESSION_TTL}


@app.post("/api/admin/logout")
def admin_logout(token: str = Depends(require_admin)):
    _sessions.pop(token, None)
    return {"success": True}


@app.get("/api/admin/prices")
def admin_prices(_: str = Depends(require_admin)):
    return {"success": True, "rows": store_db.list_price_rows()}


@app.put("/api/admin/prices/{qr_code}")
def admin_update_prices(qr_code: str, req: PriceUpdate, _: str = Depends(require_admin)):
    values, errors = store_db.validate_prices(req.model_dump())
    if errors:
        raise HTTPException(status_code=400, detail="; ".join(errors))
    product = store_db.update_prices(qr_code, values)
    if not product:
        raise HTTPException(status_code=404, detail="Unknown QR code.")
    return {"success": True, "product": product}


# Excel layout: (header, db field, editable)
EXCEL_COLUMNS = [
    ("QR Code", "qr_code", False),
    ("Product Name", "name", False),
    ("Brand", "brand", False),
    ("Store Price (₹)", "price", True),
    ("MRP (₹)", "mrp", True),
    ("Amazon Price (₹)", "amazon_price", True),
    ("Flipkart Price (₹)", "flipkart_price", True),
    ("Croma Price (₹)", "croma_price", True),
]


@app.get("/api/admin/prices.xlsx")
def admin_export_excel(_: str = Depends(require_admin)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append([h for h, _, _ in EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    locked_fill = PatternFill("solid", fgColor="E5E7EB")
    for row in store_db.list_price_rows():
        ws.append([row[field] for _, field, _ in EXCEL_COLUMNS])
    for col_idx, (_, _, editable) in enumerate(EXCEL_COLUMNS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = 45 if col_idx == 2 else 18
        if not editable:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col_idx).fill = locked_fill
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("How to use")
    for line in [
        "Edit the white price columns on the 'Prices' sheet, save, and upload the file on the admin page.",
        "Grey columns (QR Code, Product Name, Brand) are for reference; changes to them are ignored.",
        "Store Price is required. Leave a competitor price empty if that store doesn't sell the product.",
        "Rows are matched by QR Code. Unknown QR codes are reported as errors and nothing is changed.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="omnitag_prices.xlsx"'},
    )


def _match_header(text):
    """Map a spreadsheet header to a db field, tolerating small wording differences."""
    t = str(text or "").lower().replace("₹", "").replace("(", "").replace(")", "").strip()
    if t.startswith("qr"):
        return "qr_code"
    if t.startswith("store") or t in ("price", "selling price", "our price"):
        return "price"
    if t.startswith("mrp"):
        return "mrp"
    for name, col in store_db.COMPETITOR_COLUMNS.items():
        if t.startswith(name.lower()):
            return col
    return None


@app.post("/api/admin/prices/import")
async def admin_import_excel(file: UploadFile = File(...), _: str = Depends(require_admin)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx Excel file.")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File is too large (max 2 MB).")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the file. Is it a valid Excel workbook?")

    ws = wb["Prices"] if "Prices" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="The sheet is empty.")

    columns = {}
    for idx, header in enumerate(rows[0]):
        field = _match_header(header)
        if field and field not in columns:
            columns[field] = idx
    if "qr_code" not in columns or "price" not in columns:
        raise HTTPException(status_code=400, detail="The sheet needs at least 'QR Code' and 'Store Price' columns. Download the template to see the format.")

    known = store_db.existing_qr_codes()
    current = {r["qr_code"]: r for r in store_db.list_price_rows()}
    updates, errors, changes, seen = {}, [], [], set()

    for row_num, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        qr = str(row[columns["qr_code"]] or "").strip()
        if not qr:
            errors.append(f"Row {row_num}: missing QR Code")
            continue
        if qr not in known:
            errors.append(f"Row {row_num}: unknown QR Code '{qr}'")
            continue
        if qr in seen:
            errors.append(f"Row {row_num}: QR Code '{qr}' appears more than once")
            continue
        seen.add(qr)
        raw = {field: row[idx] if idx < len(row) else None for field, idx in columns.items() if field != "qr_code"}
        values, row_errors = store_db.validate_prices(raw)
        if row_errors:
            errors.extend(f"Row {row_num} ({qr}): {e}" for e in row_errors)
            continue
        diff = {k: v for k, v in values.items() if current[qr][k] != v}
        if diff:
            updates[qr] = diff
            changes.append({
                "qr_code": qr,
                "name": current[qr]["name"],
                "changes": {k: {"from": current[qr][k], "to": v} for k, v in diff.items()},
            })

    if errors:
        raise HTTPException(status_code=400, detail={"message": "No prices were changed. Fix these rows and upload again.", "errors": errors})

    store_db.bulk_update_prices(updates)
    return {"success": True, "updated": len(updates), "unchanged": len(seen) - len(updates), "changes": changes}


def sse(event, data):
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


@app.post("/api/scan_tag")
def scan_tag(req: TagRequest):
    product = resolve_product(req.qr_code, req.product)
    if not product:
        raise HTTPException(status_code=404, detail="Unknown product.")
    try:
        analysis, llm_ok, cached = assistant.get_insight(product, req.user_context)
    except llm.LLMError as e:
        return {"success": False, "error": str(e)}
    return {
        "success": True,
        "provider": f"OmniTag AI Engine ({llm.CHAT_MODEL})",
        "llm_ok": llm_ok,
        "cached": cached,
        "data": {"analysis": analysis},
    }


@app.post("/api/ask")
def ask_product(req: ChatRequest):
    product = resolve_product(req.qr_code, req.product)
    try:
        policies = rag.policy_context(req.question)
        text = assistant.answer(product, req.persona, [m.model_dump() for m in req.history], req.question, policies)
    except llm.LLMError as e:
        return {"success": False, "error": str(e)}
    return {"success": True, "answer": text, "model": llm.CHAT_MODEL, "sources": assistant.sources_for(policies)}


@app.post("/api/chat/stream")
def chat_stream(req: ChatRequest):
    """Server-Sent Events: `token` events with text pieces, then one `done` event."""
    product = resolve_product(req.qr_code, req.product)
    history = [m.model_dump() for m in req.history]

    def events():
        started = time.time()
        try:
            policies = rag.policy_context(req.question)
            for piece in assistant.stream_answer(product, req.persona, history, req.question, policies):
                yield sse("token", {"text": piece})
        except llm.LLMError as e:
            yield sse("error", {"message": str(e)})
            return
        yield sse("done", {
            "latency_ms": int((time.time() - started) * 1000),
            "model": llm.CHAT_MODEL,
            "sources": assistant.sources_for(policies),
        })

    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ---------------------------------------------------------------------------
# Website: catalogue at /, product pages at /p/<qr_code> (what QR codes link to), admin at /admin
# ---------------------------------------------------------------------------

@app.get("/p/{qr_code}")
def product_page(qr_code: str):
    return FileResponse(os.path.join(WEB_DIR, "product.html"))


@app.get("/admin")
def admin_page():
    return FileResponse(os.path.join(WEB_DIR, "admin.html"))


# Only the web/ folder is served, so the database and source files are never exposed
app.mount("/", StaticFiles(directory=WEB_DIR, html=True), name="web")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=False)
